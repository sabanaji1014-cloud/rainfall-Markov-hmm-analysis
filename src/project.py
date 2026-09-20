# NOTE: Only the input/output file paths below were updated (project.xlsx and
# the SPI_* output files now live under ../data/raw and ../data/processed,
# matching this repository's folder layout). No other logic was changed
# from the original script.

import pandas as pd
import numpy as np
import openpyxl
import scipy.stats as st

# Define SPI function
def SPI(ds, thresh):
    ds_ma = ds.rolling(thresh, center=False).mean()
    ds_ma[ds_ma <= 0] = np.nan
    ds_In = np.log(ds_ma)
    ds_In[np.isinf(ds_In)] = np.nan
    ds_mu = np.nanmean(ds_ma)
    ds_sum = np.nansum(ds_In)
    n = len(ds_In[thresh-1:])
    A = np.log(ds_mu) - (ds_sum / n)
    alpha = (1 / (4 * A)) * (1 + (1 + ((4 * A) / 3))**0.5)
    beta = ds_mu / alpha
    gamma = st.gamma.cdf(ds_ma, a=alpha, scale=beta)
    norm_spi = st.norm.ppf(gamma, loc=0, scale=1)
    return norm_spi

# Define other functions
def generate_all_patterns(pattern_length):
    num_patterns = 2**pattern_length
    all_patterns = []
    for i in range(num_patterns):
        pattern = []
        for j in range(pattern_length):
            pattern.append((i >> (pattern_length - 1 - j)) & 1)
        all_patterns.append(pattern)
    return all_patterns

def count_pattern_occurrences(data, pattern):
    pattern_count = 0
    pattern_with_rain_count = 0
    pattern_length = len(pattern)
    for i in range(len(data) - pattern_length):
        if data[i:i + pattern_length] == pattern:
            pattern_count += 1
            if data[i + pattern_length] == 1:
                pattern_with_rain_count += 1
    return pattern_count, pattern_with_rain_count

def calculate_probabilities(data, pattern_length):
    all_patterns = generate_all_patterns(pattern_length)
    pattern_probabilities = {}
    for pattern in all_patterns:
        pattern_count, rain_count = count_pattern_occurrences(data, pattern)
        total_occurrences = pattern_count
        probability_of_rain = rain_count / total_occurrences if total_occurrences > 0 else 0
        probability_of_no_rain = 1 - probability_of_rain
        pattern_probabilities[tuple(pattern)] = (probability_of_no_rain, probability_of_rain)
    return pattern_probabilities

def calculate_transition_matrix(data, pattern_length, label):
    pattern_probabilities = calculate_probabilities(data, pattern_length)
    patterns = generate_all_patterns(pattern_length)
    transition_matrix = []
    for pattern in patterns:
        no_rain_prob, rain_prob = pattern_probabilities[tuple(pattern)]
        transition_matrix.append([no_rain_prob, rain_prob])
    transition_df = pd.DataFrame(transition_matrix, index=[tuple(p) for p in patterns], columns=['P(0)', 'P(1)'])
    print(f"\n{label} Transition Matrix for Pattern Length {pattern_length}:")
    print(transition_df)
    # transition_df.to_csv(f'transition_matrix_{label}_length_{pattern_length}.csv')
    return transition_df

def categorize_spi(spi_value):
    if spi_value >= 2:
        return 'ترسالی بسیار شدید', 0
    elif 1.5 <= spi_value < 2:
        return 'ترسالی شدید', 1
    elif 1 <= spi_value < 1.5:
        return 'ترسالی متوسط', 2
    elif 0.5 <= spi_value < 1:
        return 'ترسالی ملایم', 3
    elif -0.5 <= spi_value < 0.5:
        return 'تقریباً نرمال', 4
    elif -1 <= spi_value < -0.5:
        return 'خشکسالی ملایم', 5
    elif -1.5 <= spi_value < -1:
        return 'خشکسالی متوسط', 6
    elif -2 <= spi_value < -1.5:
        return 'خشکسالی شدید', 7
    elif spi_value < -2:
        return 'خشکسالی بسیار شدید', 8
    else:
        return 'ناشناخته', 9

# Load and process data
file_path = '../data/raw/project.xlsx'
workbook = openpyxl.load_workbook(file_path)
sheet = workbook.active

data = []
for row_number in range(2, sheet.max_row + 1):
    date_cell = sheet.cell(row=row_number, column=1).value
    rain_cell = sheet.cell(row=row_number, column=6).value
    if date_cell and rain_cell:
        date = pd.to_datetime(date_cell)
        rain_value = float(rain_cell.replace(' mm', ''))
        data.append({'Date': date, 'Rain': rain_value})

df = pd.DataFrame(data)
df.set_index('Date', inplace=True)

yearly_mean_rain = df.resample('Y').mean()
spi_values_yearly = SPI(yearly_mean_rain['Rain'], 1)
spi_df_yearly = pd.DataFrame(spi_values_yearly, index=yearly_mean_rain.index, columns=['SPI'])
spi_df_yearly['SPI_binary'] = np.where(spi_df_yearly['SPI'] <= 0, 0, 1)
spi_df_yearly['Category'], spi_df_yearly['Category_code'] = zip(*spi_df_yearly['SPI'].apply(categorize_spi))
SPI_binary_yearly = spi_df_yearly['SPI_binary'].tolist()

monthly_mean_rain = df.resample('M').mean()
spi_values_monthly = SPI(monthly_mean_rain['Rain'], 1)
spi_df_monthly = pd.DataFrame(spi_values_monthly, index=monthly_mean_rain.index, columns=['SPI'])
spi_df_monthly['SPI_binary'] = np.where(spi_df_monthly['SPI'] <= 0, 0, 1)
spi_df_monthly['Category'], spi_df_monthly['Category_code'] = zip(*spi_df_monthly['SPI'].apply(categorize_spi))
SPI_binary_monthly = spi_df_monthly['SPI_binary'].tolist()

# Add mean rainfall column to monthly SPI DataFrame
spi_df_monthly['Mean_Rain'] = monthly_mean_rain['Rain'].values

# Calculate transition matrices
for pattern_length in [1, 2, 3, 4, 5]:
    calculate_transition_matrix(SPI_binary_yearly, pattern_length, 'yearly')
    calculate_transition_matrix(SPI_binary_monthly, pattern_length, 'monthly')

# Save the results to Excel files with proper date formatting
spi_df_yearly.reset_index(inplace=True)
spi_df_yearly['Date'] = spi_df_yearly['Date'].dt.strftime('%m/%d/%Y')

spi_df_monthly.reset_index(inplace=True)
spi_df_monthly['Date'] = spi_df_monthly['Date'].dt.strftime('%m/%d/%Y')

with pd.ExcelWriter('../data/processed/SPI_data_yearly_with_categories.xlsx', date_format='mm/dd/yyyy') as writer:
    spi_df_yearly.to_excel(writer, sheet_name='Yearly', index=False)

with pd.ExcelWriter('../data/processed/SPI_data_monthly_with_categories.xlsx', date_format='mm/dd/yyyy') as writer:
    spi_df_monthly.to_excel(writer, sheet_name='Monthly', index=False)
